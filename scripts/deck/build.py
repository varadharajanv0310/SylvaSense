"""
Build the deck on the organisers' own ORION_1.0_Template.pptx.

The earlier version reconstructed each slide from the exported PDF: flattened
background plates with new text boxes laid over them. This opens the real
template instead and only *adds*. Nothing the organisers drew is touched -
the vector gradients, the swoosh, the SIST badge and the section titles are
theirs and stay exactly as shipped.

Two things fall out of using the real file. The template is 20 x 11.25 in,
which is 1440 x 810 pt, identical to the PDF page, so every coordinate
carries over one to one and no font rescaling is needed. And the backgrounds
stay vector, so the deck is sharp at any projector resolution and a fraction
of the size.
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deck_content import BODY, DIAGRAM_AT, FOOTERS, CY, WH, FT, AM  # noqa: E402

TEMPLATE = r"C:\Users\varad\Downloads\ORION_1.0_Template.pptx"
BASE = r"C:\Users\varad\AppData\Local\Temp\claude\D--Sathyabama-Orion-Claude\262dadaf-978b-40de-b73d-aad2bb1e74bb\scratchpad"
DIAGRAMS = os.path.join(BASE, "diagrams")
OUT = r"D:\Sathyabama ORION Build\Order_of_One_ORION1.0_v2.pptx"

# the template is 1440 x 810 pt, so deck points are just points
K = 1.0 / 72.0          # deck pt -> inches
FONT = "Inter"          # the family the template already ships with

#: slide 1's placeholders, in the order the template lists them
TITLE_FIELDS = [
    "Problem Statement ID: ORION-PS-03",
    "Problem Statement Title: SylvaSense",
    "Team ID: [ FILL THIS IN ]",
    "Team Name: Order of One",
]


def emit(slide, d):
    """Add one line of text at deck coordinates. Never touches what exists."""
    mid = d["y"] + d["size"] * 0.62
    h = d["size"] * 2.0
    box = slide.shapes.add_textbox(
        Inches(d["x"] * K - 0.06), Inches((mid - h / 2) * K),
        Inches(min(1440 - d["x"], 1400) * K), Inches(h * K))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = d["text"]
    r.font.name = FONT
    r.font.size = Pt(round(d["size"], 1))
    r.font.bold = d["bold"]
    r.font.color.rgb = RGBColor.from_string(d["colour"].upper())


def frame(slide, name, x, y, w, h):
    """An inset diagram with a hairline, so it reads as a screen not a hole."""
    path = os.path.join(DIAGRAMS, f"{name}.png")
    if not os.path.exists(path):
        print(f"   !! missing diagram {name}.png")
        return
    pic = slide.shapes.add_picture(
        path, Inches(x * K), Inches(y * K), Inches(w * K), Inches(h * K))
    pic.line.fill.solid()
    pic.line.fill.fore_color.rgb = RGBColor.from_string("B6FAFF")
    pic.line.width = Pt(0.75)


prs = Presentation(TEMPLATE)
added = 0

# --- slide 1: fill the fields the template asks to be filled --------------
# Assigning run.text rather than text_frame.text keeps the template's own
# size, family and colour; setting text_frame.text would flatten the
# paragraph to a single unstyled run.
title_box = next(
    sh for sh in prs.slides[0].shapes
    if sh.has_text_frame and "Problem Statement" in sh.text_frame.text)
for para, value in zip(title_box.text_frame.paragraphs, TITLE_FIELDS):
    if not para.runs:
        continue
    para.runs[0].text = value
    for extra in para.runs[1:]:
        extra.text = ""          # the template splits one line across two runs
print(f"slide 1: {len(TITLE_FIELDS)} fields filled")

# the demo link is new, so it is added rather than filled
emit(prs.slides[0], dict(x=256.8, y=714.0, size=34.0, colour=CY, bold=False,
                         text="Live demo:  [ PASTE YOUR URL ]"))
added += 1

# --- slides 2-8: body, diagrams, footer, page number ---------------------
for pno in range(2, 9):
    slide = prs.slides[pno - 1]

    for d in BODY.get(pno, []):
        emit(slide, d)
        added += 1

    for (name, x, y, w, h) in DIAGRAM_AT.get(pno, []):
        frame(slide, name, x, y, w, h)
        added += 1

    if pno in FOOTERS:
        emit(slide, dict(x=76, y=763.7, size=18.0, colour=FT, bold=False,
                         text=FOOTERS[pno]))
        added += 1
    emit(slide, dict(x=1369.5, y=763.7, size=18.0, colour=FT, bold=False,
                     text=f"{pno:02d}"))
    added += 1

prs.save(OUT)
print(f"added {added} shapes across 8 slides, nothing removed")
print("wrote", OUT)
